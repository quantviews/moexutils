"""
Акции и индексы MOEX в хранилище DuckLake (lake.py), расчеты — на polars.

Таблицы:
- lake.stocks  — дневные данные акций: date, ticker, open, low, high, close,
  waprice, volume, value_rub, adj_close, shares, market_cap (ключ date + ticker);
- lake.indexes — индексы: date, ticker, BOARDID, close, value_rub, volume.

Дневные данные — из официальной истории торгов ISS (/history): close —
закрытие основной сессии, на дату — строка главного режима (максимальный
оборот). Та же методика, что у индексов.

Реестры (в metadata/, в git): сплиты, переименования, снятые с торгов,
ключевая ставка; количество акций — Excel metadata/stock-index-base.xlsx;
дивиденды — CSV соседнего проекта ../dividends/data.
"""
from __future__ import annotations

import datetime as dt
import json
import logging
import math
import os
from typing import Iterable, Optional

import numpy as np
import polars as pl
import requests

from moexutils import iss
from moexutils import lake

logger = logging.getLogger("moexutils")

# корень проекта: metadata/ и (без MOEX_DATA_ROOT) данные — на уровень выше пакета
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
METADATA_FILE = os.path.join(BASE_DIR, "metadata", "stock-index-base.xlsx")
SPLITS_FILE = os.path.join(BASE_DIR, "metadata", "splits.csv")
RENAMES_FILE = os.path.join(BASE_DIR, "metadata", "renames.csv")
DELISTED_FILE = os.path.join(BASE_DIR, "metadata", "delisted.csv")
KEY_RATE_FILE = os.path.join(BASE_DIR, "metadata", "key_rate.csv")
SECTORS_FILE = os.path.join(BASE_DIR, "metadata", "sectors.csv")
EXTERNAL_SPLITS_FILE = os.path.join(BASE_DIR, "..", "dividends", "metadata", "splits.json")
DIVIDENDS_FOLDER = os.path.join(BASE_DIR, "..", "dividends", "data")

SHARES_MARKET = 'stock/markets/shares'
INDEX_MARKET = 'stock/markets/index'
DEFAULT_INDEXES = ('IMOEX', 'MCFTR', 'RGBITR')
# Переход MOEX на расчеты T+1 по акциям; до этой даты — T+2
T1_SETTLEMENT_DATE = dt.date(2023, 7, 31)

PRICE_COLS = ('close', 'open', 'high', 'low', 'waprice')
RAW_COLS = ['date', 'ticker', 'open', 'low', 'high', 'close', 'waprice', 'volume', 'value_rub']
STOCK_COLS = RAW_COLS + ['adj_close', 'shares', 'market_cap']


def _as_date(value) -> dt.date:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    return dt.date.fromisoformat(str(value)[:10])


# ---------------------------------------------------------------- реестры metadata/

def load_splits(splits_file: Optional[str] = None, external_file: Optional[str] = None) -> pl.DataFrame:
    """
    Объединенный реестр сплитов (ticker, date, ratio, kind).

    kind: 'price' — в истории цен разрыв на дату, цены до нее делятся на ratio;
    'shares' — ISS рестейтнул цены, число акций в старых листах метаданных делится
    на ratio; 'auto' — тип по данным (есть ценовой разрыв — ценовая поправка, нет —
    поправка числа акций), ratio в ценовой семантике (дробление 1:10 → 10,
    консолидация 100:1 → 0.01). Записи внешнего реестра ../dividends (splits.json)
    получают kind='auto'; явная запись splits.csv в пределах 45 дней приоритетнее.
    """
    splits_file = splits_file or SPLITS_FILE
    external_file = external_file or EXTERNAL_SPLITS_FILE
    schema = {'ticker': pl.Utf8, 'date': pl.Date, 'ratio': pl.Float64, 'kind': pl.Utf8}
    if os.path.exists(splits_file):
        df = pl.read_csv(splits_file, try_parse_dates=True)
        if 'kind' not in df.columns:
            df = df.with_columns(pl.lit('price').alias('kind'))
        df = df.with_columns(pl.col('kind').fill_null('price'),
                             pl.col('date').cast(pl.Date), pl.col('ratio').cast(pl.Float64)
                             ).select(list(schema))
    else:
        df = pl.DataFrame(schema=schema)

    ext = []
    if os.path.exists(external_file):
        try:
            with open(external_file, encoding='utf-8') as f:
                data = json.load(f)
            for ticker, events in data.items():
                for ev in events:
                    raw = float(ev['ratio'])
                    ext.append({'ticker': ticker, 'date': _as_date(ev['date']),
                                'ratio': raw if ev.get('kind') == 'split' else 1.0 / raw, 'kind': 'auto'})
        except Exception as e:
            logger.warning(f"[WARN] Не удалось прочитать внешний реестр сплитов {external_file}: {e}")
    if ext:
        keep = [r for r in ext
                if df.filter((pl.col('ticker') == r['ticker'])
                             & ((pl.col('date') - r['date']).dt.total_days().abs() <= 45)).is_empty()]
        if keep:
            df = pl.concat([df, pl.DataFrame(keep, schema=schema)])
    return df


def load_renames(renames_file: Optional[str] = None) -> pl.DataFrame:
    """Реестр переименований (old, new, date — первый день торгов под новым тикером)."""
    path = renames_file or RENAMES_FILE
    if not os.path.exists(path):
        return pl.DataFrame(schema={'old': pl.Utf8, 'new': pl.Utf8, 'date': pl.Date})
    return pl.read_csv(path, try_parse_dates=True).with_columns(pl.col('date').cast(pl.Date))


def load_delisted(delisted_file: Optional[str] = None) -> pl.DataFrame:
    """Реестр снятых с торгов (ticker, last_date, note): история хранится, но не обновляется."""
    path = delisted_file or DELISTED_FILE
    if not os.path.exists(path):
        return pl.DataFrame(schema={'ticker': pl.Utf8, 'last_date': pl.Date, 'note': pl.Utf8})
    return pl.read_csv(path, try_parse_dates=True).with_columns(pl.col('last_date').cast(pl.Date))


def load_sectors(sectors_file: Optional[str] = None) -> pl.DataFrame:
    """Отраслевой справочник (ticker, sector); нет файла — пустая таблица."""
    path = sectors_file or SECTORS_FILE
    if not os.path.exists(path):
        return pl.DataFrame(schema={'ticker': pl.Utf8, 'sector': pl.Utf8})
    return pl.read_csv(path, schema_overrides={'ticker': pl.Utf8, 'sector': pl.Utf8}).select('ticker', 'sector')


def load_key_rate(key_rate_file: Optional[str] = None) -> pl.DataFrame:
    """История ключевой ставки ЦБ (date — дата изменения, rate — % годовых), по дате."""
    path = key_rate_file or KEY_RATE_FILE
    if not os.path.exists(path):
        return pl.DataFrame(schema={'date': pl.Date, 'rate': pl.Float64})
    return (pl.read_csv(path, try_parse_dates=True)
            .with_columns(pl.col('date').cast(pl.Date), pl.col('rate').cast(pl.Float64)).sort('date'))


def risk_free_monthly(dates, key_rate_file: Optional[str] = None) -> pl.DataFrame:
    """
    Месячная безрисковая ставка (в долях) на даты: действующая ключевая ставка / 12,
    до первой записи — первое значение. Возвращает DataFrame (date, rf).
    """
    frame = pl.DataFrame({'date': [_as_date(d) for d in dates]})
    kr = load_key_rate(key_rate_file)
    if kr.is_empty():
        return frame.with_columns(rf=pl.lit(0.0))
    first = kr['rate'][0]
    return (frame.with_row_index('_i').sort('date')
            .join_asof(kr, on='date', strategy='backward')
            .with_columns(rf=pl.col('rate').fill_null(first) / 100.0 / 12.0)
            .sort('_i').select('date', 'rf'))


CBR_KEY_RATE_URL = "https://www.cbr.ru/hd_base/KeyRate/"


def _html_table_rows(html: str) -> list[list[str]]:
    """Строки первой таблицы HTML-страницы (текст ячеек)."""
    from lxml import html as lxml_html
    tables = lxml_html.fromstring(html).xpath('//table')
    if not tables:
        return []
    return [[c.text_content().strip() for c in tr.xpath('./td|./th')] for tr in tables[0].xpath('.//tr')]


def update_key_rate(key_rate_file: Optional[str] = None, session: Optional[requests.Session] = None) -> int:
    """
    Дописывает в metadata/key_rate.csv решения ЦБ после последней записи (таблица
    ставки на cbr.ru; в файл попадают только даты изменения). Returns: число изменений.
    """
    path = key_rate_file or KEY_RATE_FILE
    kr = load_key_rate(path)
    if kr.is_empty():
        raise ValueError(f"{path}: пустой реестр — начальную историю нужно внести вручную")
    last_date, last_rate = kr['date'][-1], float(kr['rate'][-1])
    session = session or iss.make_session()
    resp = session.get(CBR_KEY_RATE_URL, headers={'User-Agent': 'Mozilla/5.0'}, params={
        'UniDbQuery.Posted': 'True', 'UniDbQuery.From': last_date.strftime('%d.%m.%Y'),
        'UniDbQuery.To': dt.date.today().strftime('%d.%m.%Y')})
    resp.raise_for_status()
    rows = []
    for cells in _html_table_rows(resp.text):
        if len(cells) < 2:
            continue
        try:
            day = dt.datetime.strptime(cells[0], '%d.%m.%Y').date()
            rate = float(cells[1].replace('\xa0', '').replace(' ', '').replace(',', '.'))
        except ValueError:
            continue  # шапка таблицы
        if day > last_date:
            rows.append((day, rate))
    changes = []
    for day, rate in sorted(rows):
        if abs(rate - last_rate) > 1e-9:
            changes.append((day, rate))
            last_rate = rate
    if not changes:
        logger.info(f"[INFO] Ключевая ставка: изменений после {last_date} нет")
        return 0
    with open(path, 'a', encoding='utf-8', newline='') as f:
        for day, rate in changes:
            f.write(f"{day:%Y-%m-%d},{rate:.2f}\n")
    logger.info(f"[OK] Ключевая ставка: +{len(changes)} изменений, последнее {changes[-1][0]} → {changes[-1][1]:.2f}%")
    return len(changes)


_shares_cache: dict = {}


def load_shares(metadata_file: Optional[str] = None) -> pl.DataFrame:
    """
    Количество акций из Excel (ticker, date, shares): листы с именами-датами
    DD.MM.YYYY, шапка на 4-й строке, колонки Code и Number of issued shares.
    Кэшируется до изменения файла.
    """
    import openpyxl

    path = metadata_file or METADATA_FILE
    schema = {'ticker': pl.Utf8, 'date': pl.Date, 'shares': pl.Float64}
    if not os.path.exists(path):
        logger.warning(f"Нет файла метаданных: {path}")
        return pl.DataFrame(schema=schema)
    key = (os.path.abspath(path), os.path.getmtime(path))
    if key in _shares_cache:
        return _shares_cache[key]
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = []
    try:
        for name in wb.sheetnames:
            try:
                day = dt.datetime.strptime(name, '%d.%m.%Y').date()
            except ValueError:
                continue
            it = wb[name].iter_rows(values_only=True)
            for _ in range(3):
                next(it, None)
            header = next(it, None)
            if not header or 'Code' not in header or 'Number of issued shares' not in header:
                continue
            ci, si = header.index('Code'), header.index('Number of issued shares')
            for r in it:
                code, n = r[ci] if ci < len(r) else None, r[si] if si < len(r) else None
                if code is None or n is None:
                    continue
                try:
                    rows.append((str(code).strip(), day, float(n)))
                except (TypeError, ValueError):
                    continue
    finally:
        wb.close()
    df = pl.DataFrame(rows, schema=schema, orient='row').sort('ticker', 'date')
    _shares_cache.clear()
    _shares_cache[key] = df
    return df


# ---------------------------------------------------------------- расчеты (polars)

def price_jump_matches(dates: list, prices: list, date, divisor: float) -> bool:
    """
    True, если в ценовом ряду на дату события разрыв, соответствующий сплиту с
    ценовым делителем divisor (история НЕ рестейтнута). Допуск — 2.5x на
    рыночное движение в день события.
    """
    date = _as_date(date)
    before = [p for d, p in zip(dates, prices) if d < date and p is not None]
    after = [p for d, p in zip(dates, prices) if d >= date and p is not None]
    if not before or not after or before[-1] <= 0 or after[0] <= 0:
        return False
    return abs(math.log((after[0] / before[-1]) / (1.0 / divisor))) < math.log(2.5)


def adjust_for_splits(df: pl.DataFrame, splits: Optional[pl.DataFrame] = None) -> pl.DataFrame:
    """
    Приводит цены (close, open, high, low, waprice) к пост-сплитовой базе по
    реестру (kind price и auto с ценовым разрывом): цены до даты делятся на ratio,
    volume умножается. adj_close, value_rub, market_cap не трогаются. Работает на
    одной бумаге и на наборе тикеров (колонка ticker).
    """
    splits = load_splits() if splits is None else splits
    splits = splits.filter(pl.col('kind').is_in(['price', 'auto']))
    if df.is_empty() or 'ticker' not in df.columns or splits.is_empty():
        return df
    df = df.sort('ticker', 'date')
    present = set(df['ticker'].unique().to_list())
    price_cols = [c for c in PRICE_COLS if c in df.columns]
    for row in splits.iter_rows(named=True):
        if row['ticker'] not in present:
            continue
        if row['kind'] == 'auto':
            one = df.filter(pl.col('ticker') == row['ticker'])
            if not price_jump_matches(one['date'].to_list(), one['close'].to_list(), row['date'], row['ratio']):
                continue
        mask = (pl.col('ticker') == row['ticker']) & (pl.col('date') < row['date'])
        exprs = [pl.when(mask).then(pl.col(c) / row['ratio']).otherwise(pl.col(c)).alias(c) for c in price_cols]
        if 'volume' in df.columns:
            exprs.append(pl.when(mask).then(pl.col('volume') * row['ratio']).otherwise(pl.col('volume')).alias('volume'))
        df = df.with_columns(exprs)
    return df


def apply_renames(df: pl.DataFrame, renames: Optional[pl.DataFrame] = None) -> pl.DataFrame:
    """
    Склеивает истории переименованных тикеров (TCSG→T, YNDX→YDEX, ...): строки
    старого тикера получают новый тикер, исходный — в колонке source_ticker;
    строки старого тикера с даты переименования отбрасываются. Цепочки A→B→C
    поддерживаются (по хронологии).
    """
    if df.is_empty() or 'ticker' not in df.columns:
        return df
    if 'source_ticker' not in df.columns:
        df = df.with_columns(pl.col('ticker').alias('source_ticker'))
    renames = load_renames() if renames is None else renames
    for row in renames.sort('date').iter_rows(named=True):
        old = pl.col('ticker') == row['old']
        dropped = df.filter(old & (pl.col('date') >= row['date'])).height
        if dropped:
            logger.warning(f"[WARN] {row['old']}: {dropped} строк с {row['date']} отброшено при склейке с {row['new']}")
            df = df.filter(~(old & (pl.col('date') >= row['date'])))
        df = df.with_columns(pl.when(pl.col('ticker') == row['old']).then(pl.lit(row['new']))
                             .otherwise(pl.col('ticker')).alias('ticker'))
    return df


def ex_dividend_pos(dates: list, record_date) -> int:
    """
    Позиция экс-дивидендной даты в торговом календаре dates (по возрастанию).
    Купивший в день D попадает в реестр, если расчеты D+k не позже даты закрытия
    реестра R: при T+1 (с 31.07.2023) экс-дата — сам R или последний торговый
    день перед R, при T+2 — торговый день перед R.
    """
    record_date = _as_date(record_date)
    k = 1 if record_date >= T1_SETTLEMENT_DATE else 2
    return int(np.searchsorted(np.array(dates, dtype='datetime64[D]'),
                               np.datetime64(record_date, 'D'), side='right')) - k


def load_dividends(ticker: str, div_folder: Optional[str] = None) -> pl.DataFrame:
    """Дивиденды тикера из CSV проекта dividends: closing_date, dividend_value (> 0), по дате."""
    path = os.path.join(div_folder or DIVIDENDS_FOLDER, f"{ticker}.csv")
    schema = {'closing_date': pl.Date, 'dividend_value': pl.Float64}
    if not os.path.exists(path):
        return pl.DataFrame(schema=schema)
    try:
        df = pl.read_csv(path, columns=['closing_date', 'dividend_value'],
                         schema_overrides={'closing_date': pl.Utf8, 'dividend_value': pl.Float64})
    except Exception as e:
        logger.error(f"{ticker}: не удалось прочитать дивиденды {path} — {e}")
        return pl.DataFrame(schema=schema)
    return (df.with_columns(pl.col('closing_date').str.to_date('%Y-%m-%d', strict=False))
            .filter(pl.col('closing_date').is_not_null() & (pl.col('dividend_value') > 0))
            .sort('closing_date'))


def adj_close(df: pl.DataFrame, dividends: Optional[pl.DataFrame] = None,
              splits: Optional[pl.DataFrame] = None) -> tuple[pl.DataFrame, list]:
    """
    adj_close одной бумаги — цена, скорректированная на дивиденды и сплиты, в
    текущей (пост-сплитовой) базе. База — сплит-скорректированный close;
    дивидендный фактор 1 − D / P по последнему закрытию с дивидендом (до экс-даты),
    P — сырая цена или сплит-скорректированная (для рестейтнутых дивидендов,
    как у ВТБ): берется первая с доходностью в (0, 50%). Объявленные дивиденды с
    отсечкой позже последней даты данных историю не корректируют.

    Returns: (df с колонкой adj_close, список отброшенных дивидендов [(date, value)]).
    """
    df = df.sort('date')
    if df.is_empty():
        return df.with_columns(pl.lit(None, pl.Float64).alias('adj_close')), []
    split_close = adjust_for_splits(df.select('date', 'ticker', 'close'), splits)['close'].to_numpy().astype(float)
    closes = df['close'].to_numpy().astype(float)
    dates = df['date'].to_list()
    adj = split_close.copy()
    skipped = []
    divs = dividends if dividends is not None else load_dividends(df['ticker'][0])
    for rec, value in reversed(list(zip(divs['closing_date'].to_list(), divs['dividend_value'].to_list()))):
        if rec > dates[-1]:
            continue  # отсечка еще не наступила
        pos = ex_dividend_pos(dates, rec)
        if pos <= 0:
            continue
        candidates = [value / b for b in (closes[pos - 1], split_close[pos - 1]) if b == b and b > 0]
        y = next((c for c in candidates if 0 < c < 0.5), None)
        if y is None:
            skipped.append((rec, value))
            continue
        adj[:pos] *= (1.0 - y)
    return df.with_columns(pl.Series('adj_close', adj, dtype=pl.Float64)), skipped


def market_cap(df: pl.DataFrame, shares: Optional[pl.DataFrame] = None,
               splits: Optional[pl.DataFrame] = None) -> pl.DataFrame:
    """
    shares и market_cap = close × shares для одной бумаги. Число акций — срезы
    Excel: между срезами действует последний, до первого — первый. Срезы до
    сплита приводятся к пост-событийной базе (kind shares; auto — если ценовой
    ряд рестейтнут). Нет данных о числе акций — колонки пустые.
    """
    df = df.sort('date')
    ticker = df['ticker'][0]
    shares = load_shares() if shares is None else shares
    known = (shares.filter(pl.col('ticker') == ticker).unique('date', keep='last').sort('date')
             .select('date', 'shares'))
    if known.is_empty():
        return df.with_columns(pl.lit(None, pl.Float64).alias('shares'),
                               pl.lit(None, pl.Float64).alias('market_cap'))
    splits = load_splits() if splits is None else splits
    for row in splits.filter((pl.col('ticker') == ticker) & pl.col('kind').is_in(['shares', 'auto'])).iter_rows(named=True):
        before = pl.col('date') < row['date']
        if row['kind'] == 'shares':
            known = known.with_columns(pl.when(before).then(pl.col('shares') / row['ratio']).otherwise(pl.col('shares')))
        elif not price_jump_matches(df['date'].to_list(), df['close'].to_list(), row['date'], row['ratio']):
            known = known.with_columns(pl.when(before).then(pl.col('shares') * row['ratio']).otherwise(pl.col('shares')))
    first = known['shares'][0]
    out = (df.drop([c for c in ('shares', 'market_cap') if c in df.columns])
           .join_asof(known, on='date', strategy='backward')
           .with_columns(pl.col('shares').fill_null(first)))
    return out.with_columns((pl.col('close') * pl.col('shares')).alias('market_cap'))


def enrich(df: pl.DataFrame, div_folder: Optional[str] = None, shares: Optional[pl.DataFrame] = None,
           splits: Optional[pl.DataFrame] = None) -> tuple[pl.DataFrame, list]:
    """adj_close и капитализация для одной бумаги; Returns: (df, отброшенные дивиденды)."""
    splits = load_splits() if splits is None else splits
    df, skipped = adj_close(df, load_dividends(df['ticker'][0], div_folder), splits)
    return market_cap(df, shares, splits), skipped


# ---------------------------------------------------------------- ISS

def fetch_stock(ticker: str, start, end=None, session: Optional[requests.Session] = None) -> pl.DataFrame:
    """
    Дневные данные акции из истории торгов ISS: на дату — строка режима с
    максимальным оборотом (главная доска), close — закрытие основной сессии.
    """
    raw = iss.security_history(SHARES_MARKET, ticker, start, end, session,
                               columns=['BOARDID', 'OPEN', 'LOW', 'HIGH', 'CLOSE', 'WAPRICE', 'VOLUME', 'VALUE'])
    if raw.is_empty():
        return pl.DataFrame(schema={c: (pl.Date if c == 'date' else pl.Utf8 if c == 'ticker' else pl.Float64)
                                    for c in RAW_COLS})
    return (raw.sort('VALUE', nulls_last=False).unique('date', keep='last')
            .filter(pl.col('CLOSE').is_not_null())
            .select(pl.col('date'), pl.lit(ticker).alias('ticker'),
                    pl.col('OPEN').alias('open'), pl.col('LOW').alias('low'), pl.col('HIGH').alias('high'),
                    pl.col('CLOSE').alias('close'), pl.col('WAPRICE').alias('waprice'),
                    pl.col('VOLUME').alias('volume'), pl.col('VALUE').alias('value_rub'))
            .sort('date'))


def fetch_index(ticker: str, start, end=None, session: Optional[requests.Session] = None) -> pl.DataFrame:
    """История индекса из ISS: date, ticker, BOARDID, close, value_rub (оборот), volume."""
    raw = iss.security_history(INDEX_MARKET, ticker, start, end, session,
                               columns=['BOARDID', 'CLOSE', 'VALUE', 'VOLUME'])
    if raw.is_empty():
        return pl.DataFrame()
    return (raw.sort('VALUE', nulls_last=False).unique('date', keep='last')
            .filter(pl.col('CLOSE').is_not_null())
            .select('date', pl.lit(ticker).alias('ticker'), 'BOARDID', pl.col('CLOSE').alias('close'),
                    pl.col('VALUE').alias('value_rub'), pl.col('VOLUME').alias('volume'))
            .sort('date'))


def is_traded(ticker: str, session: Optional[requests.Session] = None) -> Optional[bool]:
    """Торгуется ли акция хоть на одной доске рынка shares (флаг ISS is_traded); None — ISS не знает бумагу."""
    session = session or iss.make_session()
    resp = session.get(f"{iss.ISS_URL}/securities/{ticker}.json", params={'iss.only': 'boards'})
    resp.raise_for_status()
    boards = iss.to_frame(resp.json().get('boards'))
    if boards.is_empty() or 'is_traded' not in boards.columns:
        return None
    if 'market' in boards.columns:
        boards = boards.filter(pl.col('market') == 'shares')
    return bool((boards['is_traded'] == 1).any())


# ---------------------------------------------------------------- хранилище

def read_stocks(tickers=None, start=None, end=None, merge_renames: bool = True,
                split_adjusted: bool = False, columns: Optional[list[str]] = None,
                as_of: lake.AsOf = None) -> pl.DataFrame:
    """
    Дневные данные акций из хранилища. merge_renames — склейка переименованных
    тикеров (исходный тикер строки — source_ticker); split_adjusted — цены в
    пост-сплитовой базе (adjust_for_splits). as_of — данные на момент снимка
    хранилища (lake.ref); реестры сплитов и переименований — текущие.
    """
    where, params = [], []
    if tickers is not None:
        tickers = [tickers] if isinstance(tickers, str) else list(tickers)
        if merge_renames:  # склейка должна видеть и старые тикеры
            ren = load_renames()
            olds = ren.filter(pl.col('new').is_in(tickers))['old'].to_list()
            tickers = list(dict.fromkeys(tickers + olds))
        where.append(f"ticker IN ({', '.join('?' * len(tickers))})")
        params += tickers
    if start is not None:
        where.append("date >= ?")
        params.append(_as_date(start))
    if end is not None:
        where.append("date <= ?")
        params.append(_as_date(end))
    cols = ", ".join(dict.fromkeys(['date', 'ticker', *(columns or [])])) if columns else "*"
    sql = f"SELECT {cols} FROM {lake.ref('stocks', as_of)}" + (" WHERE " + " AND ".join(where) if where else "")
    df = lake.query(sql + " ORDER BY ticker, date", params)
    # Сначала склейка, потом сплиты: реестр сплитов записан на текущий тикер
    # (T, а не TCSG), и поправка должна захватить историю старого тикера
    if merge_renames:
        df = apply_renames(df)
    if split_adjusted:
        df = adjust_for_splits(df)
    return df.sort('ticker', 'date')


def list_tickers(include_delisted: bool = True, as_of: lake.AsOf = None) -> list[str]:
    """Тикеры акций в хранилище (без снятых с торгов, если include_delisted=False)."""
    names = lake.query(f"SELECT DISTINCT ticker FROM {lake.ref('stocks', as_of)} ORDER BY ticker")['ticker'].to_list()
    if not include_delisted:
        gone = set(load_delisted()['ticker'].to_list())
        names = [t for t in names if t not in gone]
    return names


def read_index(ticker: str = 'IMOEX', start=None, end=None, as_of: lake.AsOf = None) -> pl.DataFrame:
    """История индекса из хранилища: date, ticker, BOARDID, close, value_rub, volume; as_of — см. lake.ref."""
    where, params = ["ticker = ?"], [ticker.upper()]
    if start is not None:
        where.append("date >= ?")
        params.append(_as_date(start))
    if end is not None:
        where.append("date <= ?")
        params.append(_as_date(end))
    return lake.query(f"SELECT * FROM {lake.ref('indexes', as_of)} WHERE {' AND '.join(where)} ORDER BY date", params)


# сравнение «что изменилось» — общее для всех таблиц, в lake
_changed_rows = lake.changed_rows


def _final_ticker(ticker: str, renames: pl.DataFrame) -> str:
    """Итоговый тикер цепочки переименований (A→B→C: для A и B — C)."""
    step = dict(zip(renames['old'].to_list(), renames['new'].to_list()))
    seen = set()
    while ticker in step and ticker not in seen:
        seen.add(ticker)
        ticker = step[ticker]
    return ticker


def _chain_adj_close(group: pl.DataFrame, final: str, div_folder: Optional[str],
                     splits: pl.DataFrame, renames: pl.DataFrame) -> tuple[pl.DataFrame, list]:
    """
    adj_close по склеенной истории цепочки переименований (TCSG→T): поправки
    на сплиты и дивиденды итогового тикера распространяются на историю старых
    имен. Returns: (date, ticker=исходный тикер, adj_close), отброшенные дивиденды.
    """
    merged = apply_renames(group.select(RAW_COLS), renames).sort('date')
    names = merged['source_ticker'].unique().to_list()
    # дата закрытия реестра может быть в файлах двух имен; берем ее из первого
    # (итоговый тикер приоритетен), сохраняя несколько выплат на одну дату
    divs = pl.concat([load_dividends(t, div_folder).with_columns(src=pl.lit(i))
                      for i, t in enumerate(dict.fromkeys([final, *names]))])
    divs = (divs.filter(pl.col('src') == pl.col('src').min().over('closing_date'))
            .drop('src').sort('closing_date', maintain_order=True))
    out, skipped = adj_close(merged.drop('source_ticker'), divs, splits)
    return (out.select('date', 'adj_close')
            .with_columns(merged['source_ticker'].alias('ticker'))), skipped


def _recompute(frame: pl.DataFrame, div_folder: Optional[str], compute_derived: bool) -> tuple[pl.DataFrame, dict]:
    """
    adj_close и капитализация по каждому тикеру frame; Returns: (данные,
    {тикер: отброшенные дивиденды}). Для тикеров цепочки переименований adj_close
    считается по склеенной истории (в базе итогового тикера): иначе история
    старого имени не учитывала бы сплиты и дивиденды после переименования и на
    стыке был бы разрыв (TCSG→T и дробление T 1:10 в 2026 году).
    """
    shares, splits, renames = load_shares(), load_splits(), load_renames()
    parts, skipped = [], {}
    for ticker_df in frame.partition_by('ticker', maintain_order=True):
        if compute_derived:
            ticker_df, sk = enrich(ticker_df.select(RAW_COLS), div_folder, shares, splits)
            if sk:
                skipped[ticker_df['ticker'][0]] = sk
        parts.append(ticker_df.select(STOCK_COLS))
    out = pl.concat(parts, how='vertical_relaxed') if parts else frame
    if not compute_derived or out.is_empty() or renames.is_empty():
        return out, skipped

    finals = {t: _final_ticker(t, renames) for t in out['ticker'].unique().to_list()}
    for final in set(finals.values()):
        members = [t for t, f in finals.items() if f == final]
        if len(members) < 2:
            continue
        group = out.filter(pl.col('ticker').is_in(members))
        chain, sk = _chain_adj_close(group, final, div_folder, splits, renames)
        skipped.pop(final, None)
        if sk:
            skipped[final] = sk
        # строки старого имени после даты переименования в склейку не входят —
        # у них остается adj_close, посчитанный по самому тикеру
        out = (out.join(chain.rename({'adj_close': '__chain'}), on=['date', 'ticker'], how='left')
               .with_columns(pl.coalesce('__chain', 'adj_close').alias('adj_close')).drop('__chain'))
    return out.select(STOCK_COLS), skipped


def update_stocks(tickers: Optional[Iterable[str]] = None, include_delisted: bool = False,
                  div_folder: Optional[str] = None, rebuild: bool = False,
                  session: Optional[requests.Session] = None) -> int:
    """
    Дозагрузка акций из ISS в lake.stocks: с последней даты тикера (она
    перекачивается) до сегодня; rebuild — вся история с 2002 года. Для каждого
    тикера сразу пересчитываются adj_close и капитализация по всей истории, в
    хранилище пишутся только новые и изменившиеся строки — одной транзакцией.

    Returns: число записанных строк.
    """
    session = session or iss.make_session()
    existing = lake.query("SELECT * FROM lake.stocks") if 'stocks' in lake.tables() else pl.DataFrame()
    known = existing['ticker'].unique().to_list() if existing.height else []
    todo = list(tickers) if tickers is not None else known
    todo = [t.upper() for t in todo]
    if not include_delisted:
        gone = set(load_delisted()['ticker'].to_list())
        skipped = [t for t in todo if t in gone]
        todo = [t for t in todo if t not in gone]
        if skipped:
            logger.info(f"Пропущено снятых с торгов (metadata/delisted.csv): {len(skipped)}")
    logger.info(f"Акций к обновлению: {len(todo)}")

    changed, warned = [], {}
    for ticker in todo:
        old = existing.filter(pl.col('ticker') == ticker) if existing.height else pl.DataFrame()
        start = dt.date(2002, 1, 1) if rebuild or old.is_empty() else old['date'].max()
        try:
            new = fetch_stock(ticker, start, session=session)
        except Exception as e:
            logger.error(f"[ERROR] {ticker}: не удалось загрузить — {e}")
            continue
        if new.is_empty() and old.is_empty():
            logger.info(f"[SKIP] {ticker}: ISS не вернул данных")
            continue
        raw = (pl.concat([old.select(RAW_COLS), new], how='vertical_relaxed') if old.height else new)
        raw = raw.unique(['date', 'ticker'], keep='last').sort('date')
        full, sk = _recompute(raw, div_folder, compute_derived=True)
        warned.update(sk)
        delta = _changed_rows(old.select(STOCK_COLS) if old.height else pl.DataFrame(), full, ['date', 'ticker'])
        if delta.height:
            changed.append(delta)
    for ticker, items in warned.items():
        for rec, value in items:
            logger.warning(f"[WARN] {ticker}: дивиденд {value} на {rec} не согласуется ни с одной ценовой базой — пропущен")
    written = lake.write('stocks', pl.concat(changed, how='vertical_relaxed')) if changed else 0
    logger.info(f"[OK] Акции: записано строк {written}")
    return written


def recompute_stocks(tickers: Optional[Iterable[str]] = None, div_folder: Optional[str] = None,
                     dry_run: bool = False) -> pl.DataFrame:
    """
    Пересчитывает adj_close и капитализацию по всей истории (после обновления
    дивидендов, метаданных или реестра сплитов) и пишет только изменившиеся
    строки. dry_run — ничего не писать, вернуть изменившиеся строки.
    """
    frame = read_stocks(tickers, merge_renames=False)
    full, skipped = _recompute(frame, div_folder, compute_derived=True)
    for ticker, items in skipped.items():
        for rec, value in items:
            logger.warning(f"[WARN] {ticker}: дивиденд {value} на {rec} не согласуется ни с одной ценовой базой — пропущен")
    delta = _changed_rows(frame.select(STOCK_COLS), full, ['date', 'ticker'])
    if not dry_run and delta.height:
        lake.write('stocks', delta)
    logger.info(f"[OK] Пересчет adj_close и капитализации: изменилось строк {delta.height}")
    return delta


def add_stock(ticker: str, start='2002-01-01', div_folder: Optional[str] = None,
              session: Optional[requests.Session] = None) -> int:
    """Добавляет новый тикер: вся история с start, сразу с adj_close и капитализацией."""
    ticker = ticker.upper()
    new = fetch_stock(ticker, start, session=session)
    if new.is_empty():
        logger.info(f"[SKIP] {ticker}: ISS не вернул данных")
        return 0
    full, _ = _recompute(new, div_folder, compute_derived=True)
    return lake.write('stocks', full)


def update_indexes(tickers: Iterable[str] = DEFAULT_INDEXES, session: Optional[requests.Session] = None) -> int:
    """
    Дозагрузка индексов в lake.indexes с последней даты (она перекачивается);
    без истории — с 2000 года. Пишутся только новые и изменившиеся строки.
    """
    session = session or iss.make_session()
    have = (lake.query("SELECT ticker, max(date) AS last FROM lake.indexes GROUP BY ticker")
            if 'indexes' in lake.tables() else pl.DataFrame(schema={'ticker': pl.Utf8, 'last': pl.Date}))
    last = dict(zip(have['ticker'].to_list(), have['last'].to_list()))
    written = 0
    for ticker in [t.upper() for t in tickers]:
        start = last.get(ticker, dt.date(2000, 1, 1))
        try:
            new = fetch_index(ticker, start, session=session)
        except Exception as e:
            logger.warning(f"[WARN] {ticker}: не удалось обновить индекс — {e}")
            continue
        if new.is_empty():
            continue
        old = read_index(ticker, start=new['date'].min()) if ticker in last else pl.DataFrame()
        delta = _changed_rows(old.select(new.columns) if old.height else pl.DataFrame(), new, ['date', 'ticker'])
        if delta.height:
            written += lake.write('indexes', delta)
    logger.info(f"[OK] Индексы: записано строк {written}")
    return written


# ---------------------------------------------------------------- копии для SQL-потребителей
#
# Реестры metadata/ и акции со склейкой переименований и поправкой на сплиты
# лежат и в хранилище: другие проекты читают их SQL-запросом, без кода moexutils.

REGISTRIES = {
    'ref_splits': load_splits,
    'ref_renames': load_renames,
    'ref_delisted': load_delisted,
    'ref_key_rate': load_key_rate,
    'ref_sectors': load_sectors,
}


def sync_registries() -> dict[str, tuple[int, int]]:
    """Реестры metadata/ -> таблицы ref_* (только изменения); Returns: {таблица: (записано, удалено)}."""
    out = {}
    for table, loader in REGISTRIES.items():
        out[table] = lake.sync(table, loader())
        if out[table] != (0, 0):
            logger.info(f"[OK] {table}: записано {out[table][0]}, удалено {out[table][1]}")
    return out


def update_adjusted() -> tuple[int, int]:
    """
    lake.stocks_adjusted = read_stocks(split_adjusted=True): склейка переименований
    (source_ticker — исходный тикер) и цены в пост-сплитовой базе; пишутся только
    изменения. Returns: (записано, удалено).
    """
    n = lake.sync('stocks_adjusted', read_stocks(split_adjusted=True))
    logger.info(f"[OK] stocks_adjusted: записано {n[0]}, удалено {n[1]}")
    return n
