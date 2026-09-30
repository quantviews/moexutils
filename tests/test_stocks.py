"""
Тесты stocks.py: корпоративные события, adj_close, капитализация, ключевая
ставка, загрузка из ISS и хранилище. ISS замокан, хранилище — файловый
DuckLake во временной папке, реестры — временные файлы.
"""
import datetime as dt
import os

import openpyxl
import polars as pl
import pytest

from moexutils import lake
from moexutils import stocks


def d(s):
    return dt.date.fromisoformat(s)


def sdf(dates, closes, ticker='TEST', opens=None):
    n = len(dates)
    return pl.DataFrame({
        'date': [d(x) for x in dates], 'ticker': [ticker] * n,
        'open': [float(o) for o in (opens or closes)], 'low': [float(c) for c in closes],
        'high': [float(c) for c in closes], 'close': [float(c) for c in closes],
        'waprice': [float(c) for c in closes], 'volume': [10.0] * n,
        'value_rub': [float(c) * 10 for c in closes],
    })


def write_csv(path, header, rows):
    path.write_text(header + '\n' + ''.join(','.join(map(str, r)) + '\n' for r in rows), encoding='utf-8')
    return str(path)


def write_metadata(path, sheets):
    """Excel как stock-index-base.xlsx: листы-даты, шапка на 4-й строке."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for _ in range(3):
            ws.append([])
        ws.append(['Code', 'Number of issued shares'])
        for r in rows:
            ws.append(list(r))
    wb.save(path)
    return str(path)


@pytest.fixture
def registries(tmp_path, monkeypatch):
    """Пустые реестры: сплиты, внешний реестр, переименования, снятые с торгов."""
    monkeypatch.setattr(stocks, 'SPLITS_FILE', str(tmp_path / 'splits.csv'))
    monkeypatch.setattr(stocks, 'EXTERNAL_SPLITS_FILE', str(tmp_path / 'splits.json'))
    monkeypatch.setattr(stocks, 'RENAMES_FILE', str(tmp_path / 'renames.csv'))
    monkeypatch.setattr(stocks, 'DELISTED_FILE', str(tmp_path / 'delisted.csv'))
    return tmp_path


@pytest.fixture
def lake_env(tmp_path, monkeypatch):
    monkeypatch.setenv('MOEX_LAKE_CATALOG', 'ducklake:' + str(tmp_path / 'catalog.ducklake').replace('\\', '/'))
    monkeypatch.setattr(lake, 'LAKE_DATA_PATH', str(tmp_path / 'lake'))
    return tmp_path


def splits_of(tmp_path, rows):
    return stocks.load_splits(write_csv(tmp_path / 'sp.csv', 'ticker,date,ratio,kind', rows),
                              external_file=str(tmp_path / 'none.json'))


# ---------------------------------------------------------------- сплиты

class TestSplits:
    def test_prices_divided_before_split_only(self, tmp_path):
        df = sdf(['2025-01-01', '2025-01-02', '2025-01-03', '2025-01-04'], [1000, 1010, 101, 102])
        out = stocks.adjust_for_splits(df, splits_of(tmp_path, [('TEST', '2025-01-03', 10, 'price')]))
        assert out['close'].to_list() == pytest.approx([100.0, 101.0, 101.0, 102.0])
        assert out['waprice'].to_list() == pytest.approx([100.0, 101.0, 101.0, 102.0])
        assert out['volume'].to_list() == pytest.approx([100.0, 100.0, 10.0, 10.0])
        assert df['close'][0] == 1000  # исходный не изменен

    def test_reverse_split(self, tmp_path):
        out = stocks.adjust_for_splits(sdf(['2025-01-01', '2025-01-02'], [0.01, 1.0]),
                                       splits_of(tmp_path, [('TEST', '2025-01-02', 0.01, 'price')]))
        assert out['close'].to_list() == pytest.approx([1.0, 1.0])

    def test_other_tickers_and_shares_kind_untouched(self, tmp_path):
        sp = splits_of(tmp_path, [('OTHER', '2025-01-02', 10, 'price'), ('TEST', '2025-01-02', 100, 'shares')])
        out = stocks.adjust_for_splits(sdf(['2025-01-01', '2025-01-02'], [100, 101]), sp)
        assert out['close'].to_list() == [100.0, 101.0]

    def test_auto_adjusts_only_when_jump_present(self, tmp_path):
        sp = splits_of(tmp_path, [('TEST', '2025-01-03', 10, 'auto')])
        jump = stocks.adjust_for_splits(sdf(['2025-01-01', '2025-01-02', '2025-01-03'], [1000, 1010, 101]), sp)
        assert jump['close'].to_list() == pytest.approx([100.0, 101.0, 101.0])
        smooth = stocks.adjust_for_splits(sdf(['2025-01-01', '2025-01-02', '2025-01-03'], [100, 101, 102]), sp)
        assert smooth['close'].to_list() == [100.0, 101.0, 102.0]

    def test_load_splits_default_kind_and_external_json(self, tmp_path):
        csv = write_csv(tmp_path / 'sp.csv', 'ticker,date,ratio', [('AAA', '2025-01-05', 10)])
        (tmp_path / 'ext.json').write_text(
            '{"AAA": [{"date": "2025-01-02", "ratio": 10, "kind": "split"}],'
            ' "BBB": [{"date": "2025-06-01", "ratio": 100, "kind": "reverse"}]}', encoding='utf-8')
        sp = stocks.load_splits(csv, external_file=str(tmp_path / 'ext.json'))
        aaa = sp.filter(pl.col('ticker') == 'AAA')
        assert aaa.height == 1 and aaa['kind'][0] == 'price'         # явная запись побеждает (±45 дней)
        bbb = sp.filter(pl.col('ticker') == 'BBB')
        assert bbb['kind'][0] == 'auto' and bbb['ratio'][0] == pytest.approx(0.01)

    def test_real_registry(self):
        sp = stocks.load_splits(stocks.SPLITS_FILE)
        t = sp.filter(pl.col('ticker') == 'T')
        assert t.height >= 1 and t['kind'][0] == 'auto' and t['ratio'][0] == 10.0
        v = sp.filter(pl.col('ticker') == 'VTBR')
        assert v['ratio'][0] == pytest.approx(0.0002)


# ---------------------------------------------------------------- переименования

class TestRenames:
    def ren(self, tmp_path, rows):
        return stocks.load_renames(write_csv(tmp_path / 'r.csv', 'old,new,date', rows))

    def test_missing_registry_adds_source_ticker(self, tmp_path):
        out = stocks.apply_renames(sdf(['2025-01-01'], [100], 'OLD'), stocks.load_renames(str(tmp_path / 'no.csv')))
        assert out['ticker'].to_list() == ['OLD'] and out['source_ticker'].to_list() == ['OLD']

    def test_history_merged_and_old_rows_after_date_dropped(self, tmp_path):
        df = pl.concat([sdf(['2025-01-02', '2025-01-03'], [100, 999], 'OLD'),
                        sdf(['2025-01-03', '2025-01-04'], [102, 103], 'NEW')])
        out = stocks.apply_renames(df, self.ren(tmp_path, [('OLD', 'NEW', '2025-01-03')])).sort('date')
        assert set(out['ticker']) == {'NEW'}
        assert out['source_ticker'].to_list() == ['OLD', 'NEW', 'NEW']
        assert out['close'].to_list() == [100.0, 102.0, 103.0]         # мусор OLD за 03.01 отброшен

    def test_chain(self, tmp_path):
        df = pl.concat([sdf(['2025-01-01'], [1], 'A'), sdf(['2025-01-02'], [2], 'B'), sdf(['2025-01-03'], [3], 'C')])
        out = stocks.apply_renames(df, self.ren(tmp_path, [('A', 'B', '2025-01-02'), ('B', 'C', '2025-01-03')]))
        assert set(out['ticker']) == {'C'}

    def test_real_registry(self):
        r = stocks.load_renames(stocks.RENAMES_FILE)
        assert r.filter((pl.col('old') == 'TCSG') & (pl.col('new') == 'T')).height == 1


# ---------------------------------------------------------------- дивиденды и adj_close

def divs(rows):
    return pl.DataFrame({'closing_date': [d(r[0]) for r in rows], 'dividend_value': [float(r[1]) for r in rows]},
                        schema={'closing_date': pl.Date, 'dividend_value': pl.Float64})


NO_SPLITS = pl.DataFrame(schema={'ticker': pl.Utf8, 'date': pl.Date, 'ratio': pl.Float64, 'kind': pl.Utf8})


class TestExDate:
    def test_t1_record_date_is_ex_date(self):
        dates = [d('2025-07-17'), d('2025-07-18'), d('2025-07-21')]
        assert stocks.ex_dividend_pos(dates, '2025-07-18') == 1

    def test_t1_weekend_record_date(self):
        dates = [d('2025-07-17'), d('2025-07-18'), d('2025-07-21')]
        assert stocks.ex_dividend_pos(dates, '2025-07-20') == 1         # последняя торговая до R — пятница

    def test_t2_before_july_2023(self):
        dates = [d('2021-05-10'), d('2021-05-11'), d('2021-05-12')]
        assert stocks.ex_dividend_pos(dates, '2021-05-12') == 1         # экс-дата — день до отсечки


class TestAdjClose:
    def adj(self, df, rows, splits=NO_SPLITS):
        out, skipped = stocks.adj_close(df, divs(rows), splits)
        return out['adj_close'].to_list(), skipped

    def test_single_dividend(self):
        adj, _ = self.adj(sdf(['2025-01-01', '2025-01-02', '2025-01-03', '2025-01-04'], [100] * 4),
                          [('2025-01-03', 10)])
        assert adj == pytest.approx([90.0, 90.0, 100.0, 100.0])

    def test_two_dividends_compound(self):
        adj, _ = self.adj(sdf(['2025-01-01', '2025-01-02', '2025-01-03', '2025-01-04'], [100] * 4),
                          [('2025-01-02', 10), ('2025-01-03', 10)])
        assert adj == pytest.approx([81.0, 90.0, 100.0, 100.0])

    def test_t2_period(self):
        adj, _ = self.adj(sdf(['2023-05-08', '2023-05-10', '2023-05-11', '2023-05-12'], [100, 100, 95, 95]),
                          [('2023-05-11', 5)])
        assert adj == pytest.approx([95.0, 100.0, 95.0, 95.0])

    def test_large_dividend_uses_cum_price(self):
        adj, _ = self.adj(sdf(['2025-12-23', '2025-12-24', '2025-12-25', '2025-12-26'], [1799.0, 1828.2, 946.8, 961.2]),
                          [('2025-12-25', 902)])
        f = 1 - 902 / 1828.2
        assert adj == pytest.approx([1799.0 * f, 1828.2 * f, 946.8, 961.2])

    def test_before_data_future_and_none(self):
        df = sdf(['2025-01-10', '2025-01-13'], [100, 100])
        assert self.adj(df, [('2025-01-01', 10)])[0] == [100.0, 100.0]      # до истории
        assert self.adj(df, [('2025-02-20', 10)])[0] == [100.0, 100.0]      # отсечка в будущем
        assert self.adj(df, [])[0] == [100.0, 100.0]

    def test_skipped_implausible(self):
        adj, skipped = self.adj(sdf(['2025-01-01', '2025-01-02', '2025-01-03'], [100] * 3), [('2025-01-02', 80)])
        assert adj == [100.0, 100.0, 100.0] and skipped == [(d('2025-01-02'), 80.0)]

    def test_split_adjusted_base_and_dividend_across_split(self, tmp_path):
        sp = splits_of(tmp_path, [('TEST', '2025-01-03', 10, 'price')])
        df = sdf(['2025-01-01', '2025-01-02', '2025-01-03', '2025-01-04'], [1000, 1010, 101, 102])
        assert self.adj(df, [], sp)[0] == pytest.approx([100.0, 101.0, 101.0, 102.0])
        assert self.adj(df, [('2025-01-02', 100)], sp)[0] == pytest.approx([90.0, 101.0, 101.0, 102.0])

    def test_restated_dividend_basis(self, tmp_path):
        sp = splits_of(tmp_path, [('TEST', '2025-01-03', 0.0002, 'price')])
        df = sdf(['2025-01-01', '2025-01-02', '2025-01-03', '2025-01-04'], [0.02, 0.02, 100.0, 101.0])
        assert self.adj(df, [('2025-01-02', 7)], sp)[0] == pytest.approx([93.0, 100.0, 100.0, 101.0])

    def test_load_dividends(self, tmp_path):
        write_csv(tmp_path / 'TEST.csv', 'closing_date,year,period_type,dividend_value',
                  [('2025-07-18', 2024, 'full year', 34.84), ('', 2019, 'full year', 0.0), ('2020-01-01', 2019, 'x', 0)])
        out = stocks.load_dividends('TEST', str(tmp_path))
        assert out['closing_date'].to_list() == [d('2025-07-18')]
        assert stocks.load_dividends('NOPE', str(tmp_path)).is_empty()


# ---------------------------------------------------------------- капитализация

@pytest.fixture
def metadata(tmp_path):
    return write_metadata(tmp_path / 'meta.xlsx', {
        '05.01.2025': [('TEST', 1000), ('OTHER', 50)],
        '08.01.2025': [('TEST', 2000)],
        'Info': [('JUNK', 999)],
    })


class TestMarketCap:
    def test_load_shares(self, metadata):
        sh = stocks.load_shares(metadata)
        assert set(sh['date']) == {d('2025-01-05'), d('2025-01-08')} and 'JUNK' not in set(sh['ticker'])
        assert stocks.load_shares(metadata + '.none').is_empty()

    def test_load_shares_cached_until_mtime_changes(self, metadata, monkeypatch):
        stocks._shares_cache.clear()
        calls = {'n': 0}
        orig = openpyxl.load_workbook

        def counting(*a, **k):
            calls['n'] += 1
            return orig(*a, **k)

        monkeypatch.setattr(openpyxl, 'load_workbook', counting)
        stocks.load_shares(metadata)
        stocks.load_shares(metadata)
        assert calls['n'] == 1
        st = os.stat(metadata)
        os.utime(metadata, (st.st_atime, st.st_mtime + 10))
        stocks.load_shares(metadata)
        assert calls['n'] == 2

    def test_ffill_between_slices_and_first_before(self, metadata):
        out = stocks.market_cap(sdf(['2025-01-03', '2025-01-06', '2025-01-08', '2025-01-10'], [100] * 4),
                                stocks.load_shares(metadata), NO_SPLITS)
        assert out['shares'].to_list() == [1000.0, 1000.0, 2000.0, 2000.0]
        assert out['market_cap'].to_list() == [100000.0, 100000.0, 200000.0, 200000.0]

    def test_unknown_ticker_gives_nulls(self, metadata):
        out = stocks.market_cap(sdf(['2025-01-06'], [100], 'UNKNOWN'), stocks.load_shares(metadata), NO_SPLITS)
        assert out['market_cap'].to_list() == [None]

    def test_shares_kind_adjusts_old_slices(self, tmp_path):
        meta = write_metadata(tmp_path / 'm.xlsx', {'05.01.2025': [('TEST', 100000)], '08.01.2025': [('TEST', 1000)]})
        out = stocks.market_cap(sdf(['2025-01-06', '2025-01-09'], [100, 100]), stocks.load_shares(meta),
                                splits_of(tmp_path, [('TEST', '2025-01-07', 100, 'shares')]))
        assert out['market_cap'].to_list() == pytest.approx([100000.0, 100000.0])

    def test_auto_with_restated_prices_adjusts_shares(self, tmp_path):
        meta = write_metadata(tmp_path / 'm.xlsx', {'05.01.2025': [('TEST', 1000)], '08.01.2025': [('TEST', 100000)]})
        out = stocks.market_cap(sdf(['2025-01-06', '2025-01-09'], [100, 100]), stocks.load_shares(meta),
                                splits_of(tmp_path, [('TEST', '2025-01-07', 100, 'auto')]))
        assert out['market_cap'].to_list() == pytest.approx([10000000.0, 10000000.0])


# ---------------------------------------------------------------- ключевая ставка

class TestKeyRate:
    def test_risk_free_missing_file_zero(self, tmp_path):
        rf = stocks.risk_free_monthly(['2025-01-31', '2025-02-28'], str(tmp_path / 'no.csv'))
        assert rf['rf'].to_list() == [0.0, 0.0]

    def test_risk_free_ffill_and_before_start(self, tmp_path):
        path = write_csv(tmp_path / 'kr.csv', 'date,rate', [('2025-02-15', 12.0), ('2025-04-10', 24.0)])
        rf = stocks.risk_free_monthly(['2025-04-30', '2025-01-31', '2025-02-28', '2025-03-31'], path)
        assert rf['rf'].to_list() == pytest.approx([0.02, 0.01, 0.01, 0.01])   # порядок входа сохранен

    def test_real_registry_sane(self):
        kr = stocks.load_key_rate(stocks.KEY_RATE_FILE)
        assert kr.height > 50 and kr['rate'].is_between(3, 25).all()
        assert kr.filter(pl.col('date') == d('2022-02-28'))['rate'].to_list() == [20.0]

    def test_update_key_rate_appends_only_changes(self, tmp_path):
        html = """<table><tr><th>Дата</th><th>Ставка</th></tr>
        <tr><td>23.03.2026</td><td>15,00</td></tr><tr><td>20.03.2026</td><td>15,50</td></tr>
        <tr><td>16.02.2026</td><td>15,50</td></tr><tr><td>13.02.2026</td><td>16,00</td></tr></table>"""

        class S:
            def get(self, url, **kw):
                class R:
                    text = html

                    def raise_for_status(self):
                        pass
                return R()

        path = write_csv(tmp_path / 'kr.csv', 'date,rate', [('2025-12-22', '16.00')])
        assert stocks.update_key_rate(path, session=S()) == 2
        assert stocks.load_key_rate(path)['rate'].to_list() == [16.0, 15.5, 15.0]
        assert stocks.update_key_rate(path, session=S()) == 0


# ---------------------------------------------------------------- ISS

class FakeHistory:
    """ISS /history по бумаге: страницы по 2 строки."""

    def __init__(self, columns, rows):
        self.columns, self.rows = columns, rows

    def get(self, url, params=None):
        start = int((params or {}).get('start', 0))
        chunk = self.rows[start:start + 2]
        cols, total = self.columns, len(self.rows)

        class R:
            def raise_for_status(self):
                pass

            def json(self):
                return {'history': {'columns': cols, 'data': chunk},
                        'history.cursor': {'columns': ['INDEX', 'TOTAL', 'PAGESIZE'], 'data': [[start, total, 2]]}}
        return R()


def test_fetch_stock_main_board_and_drops_empty_close():
    cols = ['BOARDID', 'TRADEDATE', 'OPEN', 'LOW', 'HIGH', 'CLOSE', 'WAPRICE', 'VOLUME', 'VALUE']
    rows = [['TQBR', '2025-06-02', 99, 98, 101, 100, 99.5, 1000, 1e6],
            ['SMAL', '2025-06-02', 90, 90, 90, 90, 90, 1, 10],          # не главная доска
            ['TQBR', '2025-06-03', None, None, None, None, None, 0, 0]]   # нет торгов
    df = stocks.fetch_stock('TEST', '2025-06-01', '2025-06-03', session=FakeHistory(cols, rows))
    assert df.height == 1 and df['close'][0] == 100.0 and df['value_rub'][0] == 1e6
    assert df.columns == stocks.RAW_COLS


def test_is_traded():
    class S:
        def __init__(self, rows):
            self.rows = rows

        def get(self, url, **kw):
            rows = self.rows

            class R:
                def raise_for_status(self):
                    pass

                def json(self):
                    return {'boards': {'columns': ['boardid', 'market', 'is_traded'], 'data': rows}}
            return R()

    assert stocks.is_traded('X', S([['TQBR', 'shares', 1]])) is True
    assert stocks.is_traded('X', S([['TQBR', 'shares', 0], ['XX', 'ndm', 1]])) is False
    assert stocks.is_traded('X', S([])) is None


# ---------------------------------------------------------------- хранилище

@pytest.fixture
def store(lake_env, registries, monkeypatch):
    """Хранилище + реестры + метаданные: TEST — 1000 акций; дивиденды в папке divs."""
    meta = write_metadata(lake_env / 'meta.xlsx', {'01.01.2025': [('TEST', 1000)]})
    monkeypatch.setattr(stocks, 'METADATA_FILE', meta)
    stocks._shares_cache.clear()
    divs_dir = lake_env / 'divs'
    divs_dir.mkdir()
    return divs_dir


class TestStorage:
    def test_update_writes_derived_then_only_changes(self, store, monkeypatch):
        calls = []

        def fake_fetch(ticker, start, end=None, session=None):
            calls.append((ticker, start))
            return sdf(['2025-06-02', '2025-06-03'], [100, 100], ticker)

        monkeypatch.setattr(stocks, 'fetch_stock', fake_fetch)
        assert stocks.update_stocks(['TEST'], div_folder=str(store)) == 2
        row = stocks.read_stocks('TEST').row(0, named=True)
        assert row['adj_close'] == 100.0 and row['market_cap'] == 100000.0
        # второй прогон: с последней даты, ничего не изменилось — ничего не пишется
        assert stocks.update_stocks(div_folder=str(store)) == 0
        assert calls[-1] == ('TEST', d('2025-06-03'))

    def test_new_dividend_rewrites_history_via_recompute(self, store, monkeypatch):
        monkeypatch.setattr(stocks, 'fetch_stock',
                            lambda t, s, e=None, session=None: sdf(['2025-06-02', '2025-06-03', '2025-06-04'], [100] * 3, t))
        stocks.update_stocks(['TEST'], div_folder=str(store))
        write_csv(store / 'TEST.csv', 'closing_date,dividend_value', [('2025-06-03', 10)])
        delta = stocks.recompute_stocks(div_folder=str(store))
        assert delta.height == 1                                          # изменилась только 02.06
        assert stocks.read_stocks('TEST')['adj_close'].to_list() == pytest.approx([90.0, 100.0, 100.0])

    def test_delisted_skipped(self, store, monkeypatch):
        (store.parent / 'delisted.csv').write_text('ticker,last_date,note\nGONE,2025-01-01,x\n', encoding='utf-8')
        seen = []
        monkeypatch.setattr(stocks, 'fetch_stock',
                            lambda t, s, e=None, session=None: seen.append(t) or sdf(['2025-06-02'], [1], t))
        stocks.update_stocks(['LIVE', 'GONE'], div_folder=str(store))
        assert seen == ['LIVE']
        stocks.update_stocks(['GONE'], include_delisted=True, div_folder=str(store))
        assert seen == ['LIVE', 'GONE']

    def test_read_stocks_renames_and_splits(self, store, monkeypatch):
        (store.parent / 'renames.csv').write_text('old,new,date\nOLD,NEW,2025-06-03\n', encoding='utf-8')
        lake.write('stocks', pl.concat([sdf(['2025-06-02'], [1000], 'OLD'), sdf(['2025-06-03'], [100], 'NEW')]))
        merged = stocks.read_stocks('NEW')
        assert merged['ticker'].to_list() == ['NEW', 'NEW'] and merged['source_ticker'].to_list() == ['OLD', 'NEW']
        assert set(stocks.read_stocks(merge_renames=False)['ticker']) == {'OLD', 'NEW'}
        (store.parent / 'splits.csv').write_text('ticker,date,ratio,kind\nNEW,2025-06-03,10,price\n', encoding='utf-8')
        assert stocks.read_stocks('NEW', split_adjusted=True)['close'].to_list() == pytest.approx([100.0, 100.0])

    def test_indexes(self, store, monkeypatch):
        def fake_index(ticker, start, end=None, session=None):
            return pl.DataFrame({'date': [d('2025-06-02'), d('2025-06-03')], 'ticker': [ticker] * 2,
                                 'BOARDID': ['SNDX'] * 2, 'close': [2800.0, 2810.0],
                                 'value_rub': [1e9, 2e9], 'volume': [None, None]},
                                schema_overrides={'volume': pl.Float64})

        monkeypatch.setattr(stocks, 'fetch_index', fake_index)
        assert stocks.update_indexes(['IMOEX']) == 2
        assert stocks.update_indexes(['IMOEX']) == 0
        assert stocks.read_index('IMOEX')['close'].to_list() == [2800.0, 2810.0]


class TestRenameChainAdjClose:
    def test_old_name_history_gets_successor_split_and_dividends(self, store, monkeypatch):
        (store.parent / 'renames.csv').write_text('old,new,date\nOLD,NEW,2025-06-04\n', encoding='utf-8')
        (store.parent / 'splits.csv').write_text('ticker,date,ratio,kind\nNEW,2025-06-06,10,price\n', encoding='utf-8')
        (store / 'NEW.csv').write_text('closing_date,dividend_value\n2025-06-05,10\n', encoding='utf-8')
        rows = pl.concat([
            sdf(['2025-06-02', '2025-06-03'], [1000, 1000], 'OLD'),
            sdf(['2025-06-04', '2025-06-05', '2025-06-06', '2025-06-09'], [1000, 900, 90, 90], 'NEW')])
        full, _ = stocks._recompute(rows, str(store), compute_derived=True)
        merged = stocks.apply_renames(full).sort('date')
        # единая база NEW после дробления 1:10: дивиденд 10 при цене 1000 (фактор 0.99)
        # применен ко всем датам до экс-даты 05.06, включая историю OLD
        adj = dict(zip(merged['date'].to_list(), merged['adj_close'].to_list()))
        assert adj[d('2025-06-02')] == pytest.approx(99.0)      # OLD: /10 за сплит NEW и ×0.99 за дивиденд NEW
        assert adj[d('2025-06-03')] == pytest.approx(99.0)
        assert adj[d('2025-06-04')] == pytest.approx(99.0)
        assert adj[d('2025-06-05')] == pytest.approx(90.0)
        assert adj[d('2025-06-09')] == pytest.approx(90.0)

    def test_two_payouts_on_one_date_kept_old_name_duplicate_ignored(self, store):
        (store.parent / 'renames.csv').write_text('old,new,date\nOLD,NEW,2025-06-04\n', encoding='utf-8')
        (store.parent / 'splits.csv').write_text('ticker,date,ratio,kind\n', encoding='utf-8')
        # у NEW две выплаты с одной датой реестра (6 + 4 = 10); у OLD та же дата — дубль
        (store / 'NEW.csv').write_text('closing_date,dividend_value\n2025-06-05,6\n2025-06-05,4\n', encoding='utf-8')
        (store / 'OLD.csv').write_text('closing_date,dividend_value\n2025-06-05,10\n', encoding='utf-8')
        rows = pl.concat([
            sdf(['2025-06-02', '2025-06-03'], [1000, 1000], 'OLD'),
            sdf(['2025-06-04', '2025-06-05', '2025-06-09'], [1000, 990, 990], 'NEW')])
        full, _ = stocks._recompute(rows, str(store), compute_derived=True)
        adj = dict(zip(full['date'].to_list(), full['adj_close'].to_list()))
        assert adj[d('2025-06-02')] == pytest.approx(1000 * 0.994 * 0.996)  # обе выплаты, без дубля OLD
        assert adj[d('2025-06-04')] == pytest.approx(1000 * 0.994 * 0.996)
        assert adj[d('2025-06-09')] == pytest.approx(990.0)
